import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

def inspect_range(ws, start_row, end_row, start_col=1, end_col=15):
    print(f"\n--- {ws.title} Rows {start_row}-{end_row} Cols {start_col}-{end_col} ---")
    for r in range(start_row, end_row + 1):
        vals = []
        for c in range(start_col, end_col + 1):
            v = ws.cell(r, c).value
            if v is not None:
                vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}: {str(v)[:35]}")
        if vals:
            print(f"Row {r:3d}: " + " | ".join(vals))

print("=== INICIO ===")
inspect_range(wb['Inicio'], 1, 30, 1, 15)

print("=== PLAN BASE ===")
inspect_range(wb['Plan Base'], 1, 30, 1, 15)

print("=== FLINS SUMMARY TABLE ===")
inspect_range(wb['Flins'], 50, 70, 1, 12)

print("=== MAVUIKA SUMMARY TABLE ===")
inspect_range(wb['Mavuika'], 55, 75, 1, 12)

print("=== SANDRONE SUMMARY TABLE ===")
inspect_range(wb['Sandrone'], 1, 40, 1, 15)
inspect_range(wb['Sandrone'], 40, 80, 1, 15)
