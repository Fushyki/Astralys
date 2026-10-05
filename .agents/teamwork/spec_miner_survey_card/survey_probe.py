import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n==================== SHEET: {sheetname} (rows: {ws.max_row}, cols: {ws.max_column}) ====================")
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None and isinstance(val, str):
                val_l = val.lower().strip()
                if any(k in val_l for k in ['dps', 'dpr', 'dmgtotal', 'contrib', 'share', 'rotation', 'rot(', 'rot ']):
                    row_data = []
                    for col_idx in range(1, min(ws.max_column + 1, 30)):
                        cv = ws.cell(r, col_idx).value
                        if cv is not None:
                            col_letter = openpyxl.utils.get_column_letter(col_idx)
                            row_data.append(f"{col_letter}{r}: {str(cv)[:30]}")
                    print(f"  Matched row {r} [{val}]: {' | '.join(row_data[:8])}")
                    break
