import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

sheets_summary = {}

for s in wb.sheetnames:
    ws = wb[s]
    sheet_info = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'teams': []
    }
    
    # Check top rows (1-5) for character names, weapons, artifacts
    # Check rows 1-5 across column groups:
    # Team 1: cols B, C, D, E (2, 3, 4, 5)
    # Team 2: cols O, P, Q, R (15, 16, 17, 18)
    # Team 3: cols AB, AC, AD, AE (28, 29, 30, 31)
    # Team 4: cols AO, AP, AQ, AR (41, 42, 43, 44)
    col_groups = [
        ('Team 1', [2, 3, 4, 5]),
        ('Team 2', [15, 16, 17, 18]),
        ('Team 3', [28, 29, 30, 31]),
        ('Team 4', [41, 42, 43, 44])
    ]
    
    for tname, cgroup in col_groups:
        if cgroup[0] > ws.max_column:
            continue
        # probe row 1 to 5
        cells_1_5 = [[ws.cell(r, c).value for c in cgroup] for r in range(1, 6)]
        # probe DPR / DPS / Damage rows
        # let's find any Damage row
        dmg_row = None
        dpr_row = None
        dps_row = None
        time_row = None
        
        for r in range(1, min(ws.max_row + 1, 150)):
            val_a = str(ws.cell(r, 1).value or ws.cell(r, cgroup[0]-1).value or '').strip().lower()
            if 'damage' == val_a:
                dmg_row = r
            elif 'dmgtotal' == val_a or 'dpr' == val_a:
                dpr_row = r
            elif 'dps' in val_a:
                dps_row = r
                
        # Also check summary blocks (e.g. Flins/Mavuika row 50-65)
        summary_block = None
        for r in range(40, min(ws.max_row + 1, 140)):
            v = str(ws.cell(r, cgroup[0]).value or '').strip().lower()
            if v == 'character':
                # this is a summary block
                headers = [ws.cell(r, c).value for c in range(cgroup[0], min(cgroup[0]+6, ws.max_column+1))]
                chars_data = []
                for sub_r in range(r+1, r+5):
                    chars_data.append([ws.cell(sub_r, c).value for c in range(cgroup[0], min(cgroup[0]+6, ws.max_column+1))])
                summary_block = {
                    'header_row': r,
                    'headers': headers,
                    'rows': chars_data
                }
                break
                
        sheet_info['teams'].append({
            'team_name': tname,
            'cols': [openpyxl.utils.get_column_letter(c) for c in cgroup],
            'rows_1_to_5': cells_1_5,
            'dmg_row': dmg_row,
            'dpr_row': dpr_row,
            'dps_row': dps_row,
            'summary_block': summary_block
        })
    sheets_summary[s] = sheet_info

print(json.dumps(sheets_summary, indent=2, default=str))
