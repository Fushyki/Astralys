import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

sheets = ['Sandrone', 'Mavuika', 'Flins', 'Zibai', 'Nefer', 'Varka', 'Mualani']

for s in sheets:
    ws = wb[s]
    print(f"\n=======================================================")
    print(f"=== SUMMARY OF SHEET: {s} ===")
    print(f"=======================================================")
    
    # 1. Print row 1-5 for team 1 (B..E) and team 2 (O..R)
    for col_label, cols in [('Team 1', range(2, 6)), ('Team 2', range(15, 19))]:
        if cols.stop > ws.max_column: continue
        print(f"\n[{col_label} Header & Builds (cols {openpyxl.utils.get_column_letter(cols.start)}-{openpyxl.utils.get_column_letter(cols.stop-1)})]")
        for r in range(1, 6):
            vals = [str(ws.cell(r, c).value) for c in cols]
            print(f"  Row {r}: " + " | ".join(vals))
            
    # 2. Look for any table with 'Character'
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value).strip().lower() == 'character':
                cl = openpyxl.utils.get_column_letter(c)
                print(f"\n[Found 'Character' summary table at {cl}{r}]:")
                for sub_r in range(r, min(r + 12, ws.max_row + 1)):
                    row_data = [f"{openpyxl.utils.get_column_letter(sc)}: {ws.cell(sub_r, sc).value}" for sc in range(c, min(c + 6, ws.max_column + 1)) if ws.cell(sub_r, sc).value is not None]
                    if row_data:
                        print(f"  Row {sub_r:3d}: " + " | ".join(row_data))
                break
