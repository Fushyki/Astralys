import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

for sheetname in ['Sandrone', 'Flins', 'Nefer', 'Zibai']:
    ws = wb[sheetname]
    print(f"\n==================== {sheetname} Text cells ====================")
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val and isinstance(val, str):
                v_str = str(val).strip()
                if len(v_str) > 10 and not v_str.startswith('='):
                    # print cells with longer text (e.g. rotation strings, descriptions, notes)
                    col_letter = openpyxl.utils.get_column_letter(c)
                    print(f"  {col_letter}{r}: {v_str}")
