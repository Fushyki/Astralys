import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

sheets = ['Sandrone', 'Mavuika', 'Flins', 'Zibai', 'Nefer', 'Varka', 'Kinich', 'Mualani']

for s in sheets:
    ws = wb[s]
    print(f"\n==================== {s} FULL METRIC PROBE ====================")
    
    # 1. Damage rows
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and isinstance(v, str) and ('damage' in v.lower() or 'dpr' in v.lower() or 'dps' in v.lower() or 'dmgtotal' in v.lower()):
            cols = [f"{openpyxl.utils.get_column_letter(c)}{r}: {ws.cell(r, c).value}" for c in range(1, 10) if ws.cell(r, c).value is not None]
            print(f"  Row {r:3d} ({v}): " + " | ".join(cols[:6]))
            
    # 2. Check for rotation text or 'rot' or character action strings
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val and isinstance(val, str):
                v_str = str(val).strip()
                # Check if looks like rotation sequence or contains E, Q, N, etc.
                if any(x in v_str.lower() for x in ['rotation', 'time esped', 'combo', 'rot:']) or (' e ' in v_str.lower() and (' q ' in v_str.lower() or '>' in v_str)):
                    print(f"  Rotation-like cell {openpyxl.utils.get_column_letter(c)}{r}: {v_str[:60]}")
