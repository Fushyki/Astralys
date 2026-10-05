import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\dabiv\Downloads\Calc Sheet.xlsx', data_only=True)

sheets = ['Sandrone', 'Mavuika', 'Flins', 'Zibai', 'Nefer', 'Varka', 'Kinich', 'Mualani', 'Navia']

for s in sheets:
    ws = wb[s]
    print(f"\n==================== {s} ====================")
    for r in range(1, 10):
        row_vals = [f"{openpyxl.utils.get_column_letter(c)}{r}: {ws.cell(r, c).value}" for c in range(1, 10) if ws.cell(r, c).value is not None]
        if row_vals:
            print(f"Row {r}: " + " | ".join(row_vals))
